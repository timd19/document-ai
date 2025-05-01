import io  
import os  
from typing import Union, BinaryIO, Dict, Any  
import docx  
from pypdf import PdfReader  
import markdown  
import tempfile  

class DocumentProcessor:  
    @staticmethod  
    def extract_text_from_pdf(file_content: Union[BinaryIO, bytes]) -> str:  
        """Extract text from a PDF file."""  
        if isinstance(file_content, bytes):  
            file_content = io.BytesIO(file_content)  
        
        reader = PdfReader(file_content)  
        text = ""  
        for page in reader.pages:  
            text += page.extract_text() + "\n"  
        
        return text  
    
    @staticmethod  
    def extract_text_from_docx(file_content: Union[BinaryIO, bytes]) -> str:  
        """Extract text from a DOCX file."""  
        if isinstance(file_content, bytes):  
            file_content = io.BytesIO(file_content)  
        
        doc = docx.Document(file_content)  
        text = ""  
        for paragraph in doc.paragraphs:  
            text += paragraph.text + "\n"  
        
        return text  
    
    @staticmethod  
    def extract_text_from_markdown(file_content: Union[BinaryIO, bytes, str]) -> str:  
        """Extract text from a Markdown file."""  
        if isinstance(file_content, bytes):  
            text = file_content.decode('utf-8')  
        elif hasattr(file_content, 'read'):  
            text = file_content.read().decode('utf-8')  
        else:  
            text = file_content  
        
        # For markdown, we could either return the raw markdown or convert to HTML  
        # Here we return the raw markdown  
        return text  
    
    @staticmethod  
    def convert_markdown_to_html(markdown_text: str) -> str:  
        """Convert Markdown to HTML."""  
        return markdown.markdown(markdown_text)  
    
    @staticmethod  
    def create_pdf_from_text(text: str) -> bytes:  
        """Create a PDF from text (requires additional libraries)."""  
        # This would need a PDF generation library like reportlab  
        # For simplicity, we'll just return a placeholder  
        return f"PDF content would be generated from: {text[:100]}...".encode('utf-8')  
    
    @staticmethod  
    def create_docx_from_text(text: str) -> bytes:  
        """Create a DOCX file from text."""  
        doc = docx.Document()  
        
        # Split text by newlines and add as paragraphs  
        for paragraph in text.split('\n'):  
            if paragraph.strip():  
                doc.add_paragraph(paragraph)  
        
        # Save to a bytes buffer  
        buffer = io.BytesIO()  
        doc.save(buffer)  
        buffer.seek(0)  
        
        return buffer.read()  
    
    @staticmethod  
    def process_uploaded_file(file_content: bytes, filename: str) -> Dict[str, Any]:  
        """Process an uploaded file based on its extension."""  
        ext = os.path.splitext(filename)[1].lower()  
        
        result = {  
            "filename": filename,  
            "text": "",  
            "mime_type": "",  
        }  
        
        if ext == '.pdf':  
            result["text"] = DocumentProcessor.extract_text_from_pdf(file_content)  
            result["mime_type"] = "application/pdf"  
        elif ext == '.docx':  
            result["text"] = DocumentProcessor.extract_text_from_docx(file_content)  
            result["mime_type"] = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"  
        elif ext == '.md':  
            result["text"] = DocumentProcessor.extract_text_from_markdown(file_content)  
            result["mime_type"] = "text/markdown"  
        else:  
            # Default to treating it as plain text  
            result["text"] = file_content.decode('utf-8', errors='ignore')  
            result["mime_type"] = "text/plain"  
        
        return result